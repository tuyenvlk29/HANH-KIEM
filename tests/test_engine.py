import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))

from engine import (  # noqa: E402
    Student,
    Transaction,
    cap_bonuses,
    deduplicate,
    fold,
    match_student,
    normalize_class,
    locate_source_folder,
    iter_docx_records,
    load_students,
    load_criteria_rules,
    rule_for,
    ClassProfile,
    extract_vnedu_violation_transactions,
    is_vnedu_violation_workbook,
    save_class_template,
    vnedu_class_and_period,
    format_deduction_line,
    format_improvement_line,
    is_absence_note,
    is_uniform_only_note,
    extract_weekly_transactions,
)


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.students = [
            Student("10A1-007", 7, "Nguyễn Lê Gia Bảo", "10A1", [fold("Nguyễn Lê Gia Bảo"), fold("Gia Bảo")]),
            Student("10A1-041", 41, "Hoàng Ngọc Bảo Trang", "10A1", [fold("Hoàng Ngọc Bảo Trang"), fold("Bảo Trang")]),
        ]

    def test_class_normalization(self):
        self.assertEqual(normalize_class("10 A01"), "10A1")
        self.assertEqual(normalize_class("11-b3"), "11B3")

    def test_unique_alias(self):
        student, match_type, _ = match_student("Gia Bảo", self.students)
        self.assertEqual(student.student_id, "10A1-007")
        self.assertEqual(match_type, "unique_alias")

    def test_penalty_rules(self):
        self.assertEqual(rule_for("Không phù hệu", "TRU"), ("TRU_025_PHU_HIEU", 0.25))
        self.assertEqual(rule_for("Trễ sáng", "TRU"), ("TRU_100_DI_TRE", 1.0))
        self.assertEqual(rule_for("Vi phạm ATGT", "TRU"), ("", None))

    def test_deduction_line_uses_bullets_and_applied_points(self):
        self.assertEqual(
            format_deduction_line("Không phù hiệu 11/9 (-0,25)", 0.25),
            "- Không phù hiệu 11/9 (-0,25)",
        )
        self.assertEqual(
            format_deduction_line("Chưa làm bài tập", 0.5),
            "- Chưa làm bài tập (-0,5)",
        )
        self.assertEqual(format_improvement_line("Vắng 4 buổi có phép", 0), "- Vắng 4 buổi có phép")

    def test_only_absence_not_lateness_blocks_weekly_baseline_bonus(self):
        self.assertTrue(is_absence_note("Vắng 4b (P_3-4/7)"))
        self.assertTrue(is_absence_note("Nghỉ học có phép buổi sáng"))
        self.assertFalse(is_absence_note("Trễ sáng 15/9"))
        self.assertTrue(is_uniform_only_note("Không phù hiệu 15/9"))
        self.assertFalse(is_uniform_only_note("Không phù hiệu; Chưa làm bài tập"))

    def test_absence_withholds_all_three_weekly_baseline_items(self):
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "diem_tuan.xlsx"
            wb = Workbook()
            ws_bonus = wb.active
            ws_bonus.title = "XET TUAN 01-DIEM CONG-TONG"
            ws_deduction = wb.create_sheet("XET TUAN 01-DIEM TRU")
            ws_bonus.cell(9, 2).value = "Nguyễn Lê Gia Bảo"
            ws_bonus.cell(9, 3).value = "10A01"
            ws_bonus.cell(9, 4).value = 0.75
            ws_bonus.cell(9, 6).value = 2
            ws_bonus.cell(9, 7).value = 1
            ws_deduction.cell(9, 4).value = "Vắng học có phép ngày 8/9"
            wb.save(path)

            profile = ClassProfile("10A1", "KẾT QUẢ", 80, 10, "2026-09")
            rows, reviews = extract_weekly_transactions(path, self.students, profile)

            self.assertEqual(len(rows), 4)
            self.assertEqual(sum(row.applied_points for row in rows), 0.0)
            self.assertEqual(sum(row.proposed_points or 0 for row in rows), 0.75)
            self.assertEqual(sum(row.status == "KHONG_DU_DIEU_KIEN" for row in rows), 3)
            self.assertTrue(any(row.status == "GHI_NHAN" and "Vắng học" in row.raw_content for row in rows))
            self.assertEqual(reviews, [])

    def test_uniform_violation_keeps_study_and_attendance_but_withholds_style(self):
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "dong_phuc_tuan.xlsx"
            wb = Workbook()
            ws_bonus = wb.active
            ws_bonus.title = "XET TUAN 01-DIEM CONG-TONG"
            ws_deduction = wb.create_sheet("XET TUAN 01-DIEM TRU")
            ws_bonus.cell(9, 2).value = "Nguyễn Lê Gia Bảo"
            ws_bonus.cell(9, 3).value = "10A01"
            ws_bonus.cell(9, 4).value = 0.25
            ws_bonus.cell(9, 6).value = 1
            ws_deduction.cell(9, 4).value = "Không phù hiệu 8/9 (-0,25)"
            ws_deduction.cell(9, 5).value = -0.25
            ws_deduction.cell(9, 2).value = "Nguyễn Lê Gia Bảo"
            ws_deduction.cell(9, 3).value = "10A01"
            wb.save(path)

            profile = ClassProfile("10A1", "KẾT QUẢ", 80, 10, "2026-09")
            rows, reviews = extract_weekly_transactions(path, self.students, profile)
            points_by_code = {row.rule_code: row.applied_points for row in rows}

            self.assertEqual(points_by_code["CONG_025_HOC_TAP"], 0.25)
            self.assertEqual(points_by_code["CONG_025_CHUYEN_CAN"], 0.25)
            self.assertEqual(points_by_code["CONG_025_TAC_PHONG"], 0.0)
            self.assertEqual(points_by_code["TRU_025_PHU_HIEU"], 0.25)
            self.assertEqual(reviews, [])

    def test_bonus_cap(self):
        rows = []
        for i, value in enumerate((4.0, 4.0, 4.0), 1):
            rows.append(Transaction(str(i), "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09", "x", str(i), "a", "a", "CONG", "R", value, value, "DA_XAC_NHAN", "x", str(i)))
        cap_bonuses(rows)
        self.assertEqual(sum(x.applied_points for x in rows), 10.0)
        self.assertEqual(rows[-1].status, "VUOT_TRAN")

    def test_duplicate_same_rule_same_day(self):
        rows = [
            Transaction("1", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-18", "weekly.xlsx", "a", "Trễ sáng 18/9 (-1)", "tre sang 18 9 1", "TRU", "TRU_100_DI_TRE", 1, 1, "DA_XAC_NHAN", "a", "A"),
            Transaction("2", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-18", "Tổng hợp vi phạm tháng 09.xlsx", "b", "Trễ sáng", "tre sang", "TRU", "TRU_100_DI_TRE", 1, 1, "DA_XAC_NHAN", "b", "B"),
        ]
        result, count = deduplicate(rows)
        self.assertEqual(count, 1)
        self.assertEqual(result[1].applied_points, 0)
        self.assertEqual(result[1].status, "LOAI_BO_TRUNG_KHONG_NGAY")

    def test_undated_violation_is_removed_when_same_rule_has_a_specific_date(self):
        rows = [
            Transaction("1", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-18", "weekly.xlsx", "a", "Trễ sáng 18/9", "tre sang 18 9", "TRU", "TRU_100_DI_TRE", 1, 1, "DA_XAC_NHAN", "a", "A"),
            Transaction("2", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-30", "report.doc", "b", "Trễ sáng", "tre sang", "TRU", "TC_TRU_3_1", 1, 1, "DA_XAC_NHAN", "b", "B"),
        ]
        result, count = deduplicate(rows)
        self.assertEqual(count, 1)
        self.assertEqual(result[1].applied_points, 0)
        self.assertEqual(result[1].status, "LOAI_BO_TRUNG_KHONG_NGAY")

    def test_misspelled_uniform_violation_is_removed_when_dated_one_exists(self):
        rows = [
            Transaction("1", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-28", "weekly.xlsx", "a", "Không phù hiệu 28/9", "khong phu hieu 28 9", "TRU", "TRU_025_PHU_HIEU", 0.25, 0.25, "DA_XAC_NHAN", "a", "A"),
            Transaction("2", "10A1-007", "Nguyễn Lê Gia Bảo", "10A1", "2026-09-30", "report.xlsx", "b", "Không phù hệu", "khong phu heu", "TRU", "TC_TRU_1_1", 0.25, 0.25, "DA_XAC_NHAN", "b", "B"),
        ]
        result, count = deduplicate(rows)
        self.assertEqual(count, 1)
        self.assertEqual(result[1].applied_points, 0)
        self.assertEqual(result[1].status, "LOAI_BO_TRUNG_KHONG_NGAY")

    def test_locate_source_from_outputs_folder(self):
        located = locate_source_folder(ROOT / "outputs")
        self.assertEqual(located, (ROOT / "work" / "month_09").resolve())

    def test_word_table_reads_proposed_activity_points(self):
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "Tổ văn.docx"
            doc = Document()
            table = doc.add_table(rows=2, cols=4)
            headers = ["Họ và tên học sinh", "Lớp", "Nội dung đóng góp", "Điểm cộng đề nghị (*)"]
            values = ["Hoàng Ngọc Bảo Trang", "10A1", "Tham gia văn nghệ ngày 2/10", "1"]
            for index, value in enumerate(headers):
                table.cell(0, index).text = value
            for index, value in enumerate(values):
                table.cell(1, index).text = value
            doc.save(path)

            records = list(iter_docx_records(path, "10A1"))
            self.assertEqual(records, [(1, 2, "Hoàng Ngọc Bảo Trang", "10A1", "Tham gia văn nghệ ngày 2/10", 1.0)])

    def test_template_profile_is_inferred_from_its_own_class_and_size(self):
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "mau_lop.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "KẾT QUẢ THÁNG 10"
            ws.append(["STT", "Họ và tên học sinh", "Lớp", "Biểu hiện tích cực", "Điểm cộng", "Hành vi cần cải thiện", "Điểm trừ", "Tổng điểm", "Xếp loại", "Nhận xét"])
            ws.append([1, "Trần An", "11B03", None, None, None, None, 75, "KHÁ", None])
            ws.append([2, "Lê Bình", "11B03", None, None, None, None, 75, "KHÁ", None])
            wb.save(path)
            students, profile = load_students(path, ws.title)
            self.assertEqual(len(students), 2)
            self.assertEqual(profile.class_name, "11B3")
            self.assertEqual(profile.base_score, 75)

    def test_external_criteria_rule_overrides_generic_rule(self):
        external_rules = [("TC_TRU_2_6", 0.5, ["noi chuyen rieng trong gio hoc", "noi chuyen rieng"])]
        self.assertEqual(rule_for("Học sinh nói chuyện riêng trong giờ học", "TRU", external_rules), ("TC_TRU_2_6", 0.5))

    def test_unscored_criteria_section_does_not_inherit_previous_points(self):
        rules, report = load_criteria_rules(ROOT)
        self.assertTrue(rules)
        self.assertFalse(any(code.startswith("TC_CONG_7_") for code, _, _ in rules))
        self.assertFalse(any(code.startswith("TC_CONG_7_") for code, _, _, _ in report))

    def test_vnedu_reader_inherits_student_and_maps_short_labels(self):
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "bao_cao_vnedu.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "hs_chi_tiet"
            ws.append(["Báo cáo nề nếp"])
            ws.append([])
            ws.append(["STT", "Mã học sinh", "Họ và tên", "Lớp", "Hình thức", "Ghi chú", "Ngày"])
            ws.append([1, "HS1", "Nguyễn Lê Gia Bảo", "10A01", "Không có phù hiệu", "", "7/9/2026"])
            ws.append([None, None, None, None, "Trễ sáng", "", "8/9/2026"])
            ws.append([2, "HS2", "Hoàng Ngọc Bảo Trang", "10A01", "Nghỉ học có phép sáng (SP)", "SP", "9/9/2026"])
            wb.save(path)

            criteria = [
                ("TC_TRU_1_1", 0.25, ["khong phu hieu"]),
                ("TC_TRU_3_1", 1.0, ["di hoc tre"]),
            ]
            profile = ClassProfile("10A1", "KẾT QUẢ", 80, 10, "2026-09")
            rows, reviews = extract_vnedu_violation_transactions(path, self.students, profile, criteria)
            self.assertTrue(is_vnedu_violation_workbook(path))
            self.assertEqual(len(rows), 2)
            self.assertEqual(reviews, [])
            self.assertEqual([(x.student_name, x.event_date, x.applied_points) for x in rows], [
                ("Nguyễn Lê Gia Bảo", "2026-09-07", 0.25),
                ("Nguyễn Lê Gia Bảo", "2026-09-08", 1.0),
            ])

    def test_class_template_can_be_saved_once_for_future_vnedu_runs(self):
        with tempfile.TemporaryDirectory() as temp_name:
            root = Path(temp_name)
            template = root / "mau.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "KẾT QUẢ"
            ws.append(["STT", "Họ và tên", "Lớp", "Biểu hiện tích cực", "Điểm cộng", "Hành vi cần cải thiện", "Điểm trừ", "Tổng điểm", "Xếp loại", "Nhận xét"])
            ws.append([1, "Trần An", "11B03", None, None, None, None, 80, "TỐT", None])
            wb.save(template)

            saved, profile = save_class_template(template, root)
            self.assertTrue(saved.exists())
            self.assertEqual(profile.class_name, "11B3")

            report = root / "vnedu.xlsx"
            report_wb = Workbook()
            report_ws = report_wb.active
            report_ws.append(["STT", "Mã", "Họ và tên", "Lớp", "Hình thức", "Ghi chú", "Ngày"])
            report_ws.append([1, "HS1", "Trần An", "11B03", "Trễ sáng", "", "2/10/2026"])
            report_wb.save(report)
            self.assertEqual(vnedu_class_and_period(report), ("11B3", "2026-10"))


if __name__ == "__main__":
    unittest.main()
